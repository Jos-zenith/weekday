"""
Builds Weekday_Scheduling_Desk.xlsx - the interns' scheduling sheet (also loads in Google Sheets).

Design rules (from review):
  * Interns see four tabs: START HERE, RAM, SHYAM, LOG. Everything else is the lead's.
  * One LOG line per conversation. A second line is only needed when the answer differs by company.
  * Messages never promise a held slot. We book for a candidate only where the partner said YES in PARTNERS;
    otherwise we send the partner's own booking link. Slots are picked live in the partner's calendar.
  * Every WhatsApp message is short enough for a pre-filled link (Excel's HYPERLINK limit is 255 characters).

All logic is live formulas driven by INPUT (partner data), PARTNERS and LOG. Demo data is SYNTHETIC
(example.com addresses, fake phone numbers), generated from the analysis model so the dry run has known truth.
"""
import datetime as dt
import json
import os
import sys

import numpy as np
import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "analysis"))
import sim2  # noqa: E402

OUT = os.environ.get("DESK_OUT", os.path.join(ROOT, "Weekday_Scheduling_Desk.xlsx"))
TRUTH = os.path.join(ROOT, "tests", "output", "truth.json")
NROWS = 600            # INPUT / DESK capacity (rows 2..601)
LOGN = 2000            # LOG capacity (rows 2..2001); one sheet per daily batch
NQ = 25                # rows shown on each intern's list
DAY0 = dt.datetime(2026, 10, 13)

F = "Arial"
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_title = Font(name=F, size=14, bold=True)
f_big = Font(name=F, size=20, bold=True)
f_input = Font(name=F, size=10, color="0000FF")
f_link = Font(name=F, size=10, color="0563C1", underline="single")
f_note = Font(name=F, size=9, italic=True, color="595959")
fill_head = PatternFill("solid", fgColor="1F3864")
fill_in = PatternFill("solid", fgColor="FFF2CC")
fill_sec = PatternFill("solid", fgColor="D9E1F2")
fill_ok = PatternFill("solid", fgColor="E2EFDA")
fill_warn = PatternFill("solid", fgColor="FCE4D6")
fill_urgent = PatternFill("solid", fgColor="FFE699")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")


def header(ws, row, labels, col=1, widths=None):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t)
        c.font, c.fill, c.border = f_head, fill_head, box
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w_ in enumerate(widths):
            ws.column_dimensions[get_column_letter(col + i)].width = w_


def section(ws, row, text, ncols=6):
    ws.cell(row, 1, text).font = f_bold
    for j in range(1, ncols + 1):
        ws.cell(row, j).fill = fill_sec


# ----------------------------------------------------------------------------------------------
# 1. synthetic demo data from the analysis model (so the dry run has known truth)
# ----------------------------------------------------------------------------------------------
cfg = {**sim2.BASE}
world = sim2.make_world(np.random.default_rng(2026), cfg)
rng = np.random.default_rng(7)
n, nc = world["n"], world["nc"]
FIRST = ["Aarav", "Diya", "Kabir", "Ananya", "Vihaan", "Ira", "Reyansh", "Meera", "Arjun", "Saanvi", "Rohan",
         "Kiara", "Dev", "Navya", "Ishaan", "Tara", "Yash", "Zoya", "Neil", "Riya", "Karan", "Pooja", "Aditya",
         "Sneha", "Manav", "Isha", "Rahul", "Nisha", "Varun", "Kavya"]
LASTI = list("ABCDEFGHJKLMNPRSTV")
cand_name, cand_email, cand_phone = [], [], []
for c in range(nc):
    fn = FIRST[c % len(FIRST)]
    ln = LASTI[(c // len(FIRST)) % len(LASTI)]
    cand_name.append(f"{fn} {ln}.")
    em = f"{fn.lower()}.{ln.lower()}{c:03d}@example.com"
    ph = str(9000000000 + c * 13 + 7)
    if not world["valid"][c]:                      # unusable details: bad email OR short phone
        if c % 2:
            em = em.replace("@", "")
        else:
            ph = ph[:-1]
    cand_email.append(em)
    cand_phone.append(ph)

rows = []
for i in range(n):
    c, comp, rd = int(world["cand_of"][i]), int(world["comp"][i]), int(world["rnd"][i]) + 1
    rows.append(dict(company=f"Partner {comp:03d}", intv_name=f"Interviewer {(c + comp) % 3 + 1}",
                     intv=f"interviewer{(c + comp) % 3}@partner{comp:03d}.example.com",
                     cand=cand_name[c], email=cand_email[c], phone=cand_phone[c], round=f"Round{rd}",
                     method=f"https://calendly.com/partner{comp:03d}/r{rd}",
                     added=DAY0.replace(hour=19, minute=51), signal=["Clicked", "Opened", "None"][int(world["sig"][c])],
                     c=c, comp=comp, i=i))
for i in rng.choice(n, size=int(round(0.04 * n)), replace=False):      # partners re-sending the same record
    d = dict(rows[int(i)])
    d["added"] = d["added"] + dt.timedelta(minutes=int(rng.integers(1, 6)))
    d["dup_of"] = int(i)
    rows.append(d)
rows = [rows[int(k)] for k in rng.permutation(len(rows))]
assert len(rows) <= NROWS

comps = sorted({r["company"] for r in rows})
consent = {cn: ("YES" if rng.random() < 0.7 else "NO") for cn in comps}     # demo: 70% of partners allow booking for them

os.makedirs(os.path.dirname(TRUTH), exist_ok=True)
json.dump(dict(
    cand={cand_phone[c]: dict(cls=int(world["cls"][c]), valid=bool(world["valid"][c])) for c in range(nc)},
    req=[dict(phone=r["phone"], company=r["company"], round=r["round"],
              self_min=(None if not np.isfinite(world["self_time"][r["i"]]) else float(world["self_time"][r["i"]])))
         for r in rows if "dup_of" not in r],
    consent=consent, day0=DAY0.isoformat(), arrival_min=int(cfg["arrive_clock"]),
), open(TRUTH, "w"))

# planning table: chance a try ends in a booking, by engagement signal x try number (only used to ORDER the list)
prior, sig_m = np.array(cfg["prior"]), np.array(cfg["sig"])
ans, conv = np.array(cfg["ans_p"]), cfg["conv_p"]
ptab = []
for s in range(3):
    post = prior * sig_m[:, s]
    post /= post.sum()
    r_ = []
    for j in range(3):
        r_.append(float(round((post * ans).sum() * conv, 2)))
        post = post * (1 - ans)
        post /= post.sum()
    ptab.append(r_)

# ----------------------------------------------------------------------------------------------
wb = openpyxl.Workbook()
wb.remove(wb.active)
names = ["START HERE", "RAM", "SHYAM", "LOG", "DASHBOARD", "PARTNERS", "INPUT", "SETTINGS", "DESK", "LISTS"]
S = {k: wb.create_sheet(k) for k in names}

# ---------------------------------------------------------------------------------- SETTINGS (lead)
ws = S["SETTINGS"]
ws["A1"], ws["A1"].font = "SETTINGS (lead only). Blue cells can be changed.", f_title
header(ws, 2, ["Setting", "Value", "Notes"], widths=[52, 18, 96])
settings = [
    # key, row, label, value, number format, note
    ("testclock", 3, "Test clock (LIVE USE: leave blank)", dt.datetime(2026, 10, 14, 10, 0), "yyyy-mm-dd hh:mm",
     "Only for the demo and tests. Blank = the sheet uses the real time."),
    ("clock", 4, "Clock used by the sheet", '=IF(B3<>"",B3,NOW())', "yyyy-mm-dd hh:mm", "Do not edit."),
    ("sla", 5, "Deadline: hours after the request arrives", 24, "0", "The brief: scheduled within 24 hours."),
    ("maxtry", 6, "Most tries per person", 3, "0", "After this many tries the person leaves the list."),
    ("win0", 7, "No contact before", dt.time(10, 0), "hh:mm", ""),
    ("win1", 8, "No contact after", dt.time(21, 0), "hh:mm", ""),
    ("ram0", 9, "Ram shift starts", dt.time(10, 0), "hh:mm", "While both are on shift, people are split between them by phone number."),
    ("ram1", 10, "Ram shift ends", dt.time(18, 0), "hh:mm", ""),
    ("shy0", 11, "Shyam shift starts", dt.time(12, 0), "hh:mm", "One intern off sick: set their start and end to the same time."),
    ("shy1", 12, "Shyam shift ends", dt.time(20, 0), "hh:mm", ""),
    ("mintry", 13, "Minutes per try (call + WhatsApp)", 4.5, "0.0", "Time 20 real tries on day 1 and put the average here."),
    ("urgent", 14, "Urgent when hours left are at most", 4, "0", "Urgent rows turn yellow and move up."),
    ("ubst", 15, "Urgent priority boost", 2, "0.0", ""),
    ("rbst", 16, "Round 2 / Round 3 priority boost", 1.3, "0.0", "Later rounds are warmer and block the partner's pipeline."),
    ("target", 17, "Target: share scheduled within the deadline", 0.8, "0%", "The brief's 80%."),
    ("pbook", 18, "Planning value: requests booked per try", 0.5, "0.00",
     "ASSUMPTION. Replaced automatically by today's measured value once 30 tries are logged."),
    ("pself", 19, "Planning value: share of requests candidates book themselves", 0.25, "0%",
     "ASSUMPTION. Replace with yesterday's measured value (DASHBOARD row 11 / row 6)."),
    ("stip", 21, "Cost of one intern per month (Rs)", 20000, "#,##0",
     "Advertised founder's-office intern stipends in Bengaluru cluster around Rs 15-30k."),
    ("days", 22, "Working days per month", 22, "0", ""),
    ("fee", 23, "Weekday fee per hire (share of annual CTC)", 0.15, "0%", "Weekday's published contingency model is 15% of annual salary."),
    ("ctc", 24, "Typical annual CTC of a placed candidate (Rs)", 1800000, "#,##0", "ASSUMPTION. Use Weekday's real average."),
    ("hire", 25, "Hires per scheduled interview", 0.03, "0.0%", "ASSUMPTION (1 hire per ~33 interviews). Use Weekday's funnel data."),
    ("lost", 26, "Share of requests that miss 24h and are then never interviewed", 0.25, "0%",
     "ASSUMPTION. Measure: of requests that missed 24h last week, how many were never interviewed?"),
    ("wamsg", 27, "WhatsApp Business API cost per template message (Rs)", 0.86, "0.00",
     "Meta's 2026 India marketing rate, before GST and provider fees. Utility rate is about Rs 0.12."),
]
SET = {}
for key, r, lab, val, fmt, note in settings:
    ws.cell(r, 1, lab).font = f_norm
    c = ws.cell(r, 2, val)
    c.number_format = fmt
    if key == "clock":
        c.font = f_norm
    else:
        c.font, c.fill = f_input, fill_in
    ws.cell(r, 3, note).font = f_norm
    ws.cell(r, 3).alignment = wrap
    SET[key] = f"SETTINGS!$B${r}"
section(ws, 20, "Cost inputs (used on DASHBOARD)", 3)
header(ws, 29, ["Chance a try ends in a booking (planning values, used only to ORDER the list)", "Try 1", "Try 2", "Try 3"])
for k, (lab, vals) in enumerate(zip(["Clicked", "Opened", "None"], ptab)):
    ws.cell(30 + k, 1, lab).font = f_norm
    for j, v in enumerate(vals):
        c = ws.cell(30 + k, 2 + j, v)
        c.font, c.fill, c.number_format = f_input, fill_in, "0.00"
ws["A33"] = ("Signal = what the platform's first-touch report shows for the person (clicked the link, opened the message, or "
             "nothing). Blank counts as None. The values come from the analysis model and only change the order of the list.")
ws["A33"].font, ws["A33"].alignment = f_note, wrap
ws.merge_cells("A33:C33")
ws.row_dimensions[33].height = 30
ws.column_dimensions["D"].width = 10
CK = SET["clock"]

# ---------------------------------------------------------------------------------- LISTS (hidden)
ws = S["LISTS"]
header(ws, 1, ["Outcome", "When to use it", "Counts as a try (1/0)", "Wait before next try (min)", "Scope"],
       widths=[18, 70, 12, 14, 30])
outcomes = [
    ("NO_ANSWER", "Called, no answer. You sent the WhatsApp.", 1, 120, "Person"),
    ("CALL_BACK", "They picked up but asked you to call later.", 1, 120, "Person"),
    ("BOOKED", "You booked the interview for them in the partner's calendar.", 1, 0, "Company in E, or blank = all"),
    ("SENT_LINK", "They said yes; you sent the partner's booking link (partner does not let us book).", 1, 180, "Person"),
    ("NOT_INTERESTED", "They no longer want this interview.", 1, 0, "Company in E, or blank = all"),
    ("STOP", "They asked us to stop contacting them.", 1, 0, "Person"),
    ("WRONG_NUMBER", "Wrong number, or not their WhatsApp.", 1, 0, "Person"),
    ("SELF_BOOKED", "Their booking is already in the partner's calendar (they did it).", 0, 0, "Company in E, or blank = all"),
    ("NO_SLOTS", "The partner's calendar has no free time. Write the company in E.", 0, 0, "Company"),
]
for i, row in enumerate(outcomes, start=2):
    for j, v in enumerate(row, start=1):
        c = ws.cell(i, j, v)
        c.font, c.border, c.alignment = f_norm, box, wrap
NOUT = len(outcomes)
OUTR = f"LISTS!$A$2:$E${NOUT + 1}"
# message parts: opening by situation + ending by whether we may book for them
header(ws, 13, ["Part", "Text ({NAME}, {INTERN}, {COMPANY}, {LINK} are filled in automatically)"])
parts = [
    ("T1", "Hi {NAME}, it's {INTERN} from Weekday. You have an interview request from {COMPANY}. "),
    ("T2", "Hi {NAME}, {INTERN} from Weekday again about your {COMPANY} interview. "),
    ("T3", "Hi {NAME}, last reminder from Weekday about your {COMPANY} interview. "),
    ("CHK", "Hi {NAME}, your {COMPANY} interview isn't booked yet. "),
    ("YES", "Reply YES and I'll book the first free slot."),
    ("NO", "Pick a time: {LINK}"),
]
# Kept short on purpose: with a long company name and booking link the pre-filled WhatsApp link is still
# under Excel's 255-character HYPERLINK limit (tests/check_workbook.py measures it).
for i, (k, t) in enumerate(parts, start=14):
    ws.cell(i, 1, k).font = f_bold
    ws.cell(i, 2, t).font = f_norm
PARTR = "LISTS!$A$14:$B$19"

# ---------------------------------------------------------------------------------- INPUT (lead)
ws = S["INPUT"]
header(ws, 1, ["Company", "Interviewer", "Interviewer Email", "Candidate", "Candidate Email", "Candidate Phone",
               "Scheduling method", "Round", "Added On", "Signal (optional: Clicked / Opened / None)"],
       widths=[14, 14, 32, 14, 32, 14, 38, 9, 16, 16])
for k, r in enumerate(rows, start=2):
    vals = [r["company"], r["intv_name"], r["intv"], r["cand"], r["email"], r["phone"], r["method"], r["round"],
            r["added"], r["signal"]]
    for j, v in enumerate(vals, start=1):
        ws.cell(k, j, v).font = f_input
    ws.cell(k, 6).number_format = "@"
    ws.cell(k, 9).number_format = "yyyy-mm-dd hh:mm"
ws["L1"] = ("PASTE the partner sheet into columns A-I from row 2 (same 9 columns as the sheet Weekday sends). "
            "Column J is optional. Demo rows below are SYNTHETIC (example.com addresses, fake numbers).")
ws["L1"].font, ws["L1"].alignment = f_bold, wrap
ws.column_dimensions["L"].width = 60
dv = DataValidation(type="list", formula1='"Clicked,Opened,None"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"J2:J{NROWS + 1}")
ws.freeze_panes = "A2"

# ---------------------------------------------------------------------------------- PARTNERS (lead)
ws = S["PARTNERS"]
header(ws, 1, ["Company", "Booking link (their calendar)", "We may book for candidates? (YES / NO)",
               "Slots (OK / NO SLOTS / REOPENED)", "Partner contact", "Requests today", "Scheduled in 24h", "Rate",
               "Interns reported NO_SLOTS"],
       widths=[14, 40, 16, 16, 30, 11, 11, 8, 13])
R = NROWS + 1
for k, cname in enumerate(comps, start=2):
    slug = cname.lower().replace(" ", "")
    ws.cell(k, 1, cname).font = f_norm
    ws.cell(k, 2, f"https://calendly.com/{slug}/interview").font = f_input
    c = ws.cell(k, 3, consent[cname]); c.font, c.fill = f_input, fill_in
    c = ws.cell(k, 4, "OK"); c.font, c.fill = f_input, fill_in
    ws.cell(k, 5, f"ops@{slug}.example.com").font = f_input
    ws.cell(k, 6, f'=COUNTIFS(DESK!$B$2:$B${R},A{k},DESK!$L$2:$L${R},1)').font = f_norm
    ws.cell(k, 7, f'=COUNTIFS(DESK!$B$2:$B${R},A{k},DESK!$AO$2:$AO${R},1)').font = f_norm
    ws.cell(k, 8, f'=IF(F{k}=0,"",G{k}/F{k})').number_format = "0%"
    ws.cell(k, 9, f'=COUNTIFS(LOG!$E$2:$E${LOGN+1},A{k},LOG!$D$2:$D${LOGN+1},"NO_SLOTS")').font = f_norm
NPART = len(comps)
ws.cell(NPART + 3, 1, ("Column C: put YES only when the partner has agreed in writing that we may book interviews in their calendar "
                       "for the candidate. Blank or NO = we send their link instead. Column D: if an intern logs NO_SLOTS the "
                       "partner's requests pause; email the partner, and type REOPENED once they add slots.")).font = f_note
ws.cell(NPART + 3, 1).alignment = wrap
ws.merge_cells(start_row=NPART + 3, start_column=1, end_row=NPART + 3, end_column=9)
ws.row_dimensions[NPART + 3].height = 40
dv = DataValidation(type="list", formula1='"YES,NO"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"C2:C{NPART + 1}")
dv = DataValidation(type="list", formula1='"OK,NO SLOTS,REOPENED"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"D2:D{NPART + 1}")
ws.freeze_panes = "B2"
PA = f"PARTNERS!$A$2:$A${NPART + 1}"
PCOL = lambda c: f"PARTNERS!${c}$2:${c}${NPART + 1}"

# ---------------------------------------------------------------------------------- LOG (interns)
ws = S["LOG"]
header(ws, 1, ["Time (Ctrl+Shift+;)", "Who", "Phone (copy from your list)", "Outcome", "Company (blank = all of theirs)",
               "Notes", "time used (auto)", "try? (auto)", "try key (auto)", "Check (auto)"],
       widths=[18, 10, 16, 18, 22, 40, 16, 8, 16, 34])
L = lambda c: f"LOG!${c}$2:${c}${LOGN + 1}"
for r in range(2, LOGN + 2):
    ws.cell(r, 1).number_format = "yyyy-mm-dd hh:mm"
    ws.cell(r, 3).number_format = "@"
    for c in range(1, 7):
        ws.cell(r, c).font = f_input
    # G: a time typed alone (e.g. 14:35) takes the date of the latest line above; rolls to the next day if earlier
    if r == 2:
        g = f'=IF(A2="","",IF(A2>=1,A2,INT({CK})+A2))'
        h = f'=IF(D2="","",IFERROR(VLOOKUP(D2,{OUTR},3,FALSE),0))'
    else:
        prev = f"MAX($G$2:G{r-1})"
        g = (f'=IF(A{r}="","",IF(A{r}>=1,A{r},IF({prev}=0,INT({CK})+A{r},'
             f'INT({prev})+A{r}+IF(A{r}<MOD({prev},1),1,0))))')
        # a second line about the same person within 15 minutes is the same conversation, not a new try
        h = (f'=IF(D{r}="","",IF(IFERROR(VLOOKUP(D{r},{OUTR},3,FALSE),0)=0,0,'
             f'IF(IFERROR(COUNTIFS($C$2:C{r-1},C{r},$H$2:H{r-1},1,$G$2:G{r-1},">="&(G{r}-15/1440)),0)>0,0,1)))')
    ws.cell(r, 7, g).number_format = "yyyy-mm-dd hh:mm"
    ws.cell(r, 8, h)
    ws.cell(r, 9, f'=IF(H{r}=1,C{r}&"|"&COUNTIFS($C$2:C{r},C{r},$H$2:H{r},1),"")')
    ws.cell(r, 10, (f'=IF(D{r}="","",IF(A{r}="","Missing time",IF(D{r}="NO_SLOTS",'
                    f'IF(COUNTIF({PA},E{r})=0,"Write the company in column E","OK"),'
                    f'IF(C{r}="","Missing phone",IF(COUNTIF(DESK!$F$2:$F${R},C{r})=0,"Phone not found: copy it from your list",'
                    f'IF(AND(E{r}<>"",COUNTIFS(DESK!$F$2:$F${R},C{r},DESK!$B$2:$B${R},E{r})=0),"Company is not one of theirs","OK"))))))'))
    for c in range(7, 11):
        ws.cell(r, c).font = f_note
ws.conditional_formatting.add(f"J2:J{LOGN + 1}", FormulaRule(formula=['AND($J2<>"",$J2<>"OK")'], fill=fill_warn))
dv = DataValidation(type="list", formula1=f"=LISTS!$A$2:$A${NOUT + 1}", allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"D2:D{LOGN + 1}")
dv2 = DataValidation(type="list", formula1='"Ram,Shyam,Lead"', allow_blank=True)
ws.add_data_validation(dv2)
dv2.add(f"B2:B{LOGN + 1}")
dv3 = DataValidation(type="list", formula1=f"={PA}", allow_blank=True)
ws.add_data_validation(dv3)
dv3.add(f"E2:E{LOGN + 1}")
ws.freeze_panes = "A2"
LT, LP, LO, LC, LTRY, LKEY = L("G"), L("C"), L("D"), L("E"), L("H"), L("I")

# ---------------------------------------------------------------------------------- DESK (hidden engine)
ws = S["DESK"]
cols = [
    ("A", "#"), ("B", "Company"), ("C", "Interviewer email"), ("D", "Candidate"), ("E", "Candidate email"), ("F", "Phone"),
    ("G", "Round"), ("H", "Added on"), ("I", "Signal"), ("J", "Request key"), ("K", "Record check"), ("L", "Unique?"),
    ("M", "Booked?"), ("N", "Declined?"), ("O", "Person stopped?"), ("P", "Partner blocked?"), ("Q", "Tries"),
    ("R", "Last try"), ("S", "Last outcome"), ("T", "Wait (min)"), ("U", "Next allowed"), ("V", "Deadline"),
    ("W", "Hours left"), ("X", "Open?"), ("Y", "Due now?"), ("Z", "Owner"), ("AA", "Open requests (person)"),
    ("AB", "Seq"), ("AC", "May book for them?"), ("AD", "Booking link"), ("AE", "P(try books)"), ("AF", "Round no."),
    ("AG", "Score"), ("AH", "Rank"), ("AI", "Queue key"), ("AJ", "Seq key"), ("AK", "Person first row?"),
    ("AL", "State"), ("AM", "Booked at"), ("AN", "Hours to schedule"), ("AO", "Scheduled in 24h?"),
    ("AP", "Booked by us?"), ("AQ", "First name"),
]
header(ws, 1, [t for _, t in cols])
for col, _ in cols:
    ws.column_dimensions[col].width = 13
ws.column_dimensions["J"].width = 30
ws.column_dimensions["AL"].width = 28
D_ = lambda c: f"${c}$2:${c}${R}"
TOD = f"ROUND(MOD({CK},1)*1440,0)"
mins = lambda key: f"ROUND({SET[key]}*1440,0)"
for r in range(2, R + 1):
    w_ = {}
    w_["A"] = r - 1
    w_["B"] = f'=IF(INPUT!A{r}="","",TRIM(INPUT!A{r}))'
    w_["C"] = f'=IF(B{r}="","",INPUT!C{r})'
    w_["D"] = f'=IF(B{r}="","",INPUT!D{r})'
    w_["E"] = f'=IF(B{r}="","",INPUT!E{r})'
    w_["F"] = f'=IF(B{r}="","",TRIM(INPUT!F{r}&""))'
    w_["G"] = f'=IF(B{r}="","",INPUT!H{r})'
    w_["H"] = f'=IF(B{r}="","",INPUT!I{r})'
    w_["I"] = f'=IF(B{r}="","",IF(INPUT!J{r}="","None",INPUT!J{r}))'
    w_["J"] = f'=IF(B{r}="","",LOWER(E{r})&"|"&LOWER(B{r})&"|"&G{r}&"|"&LOWER(C{r}))'
    # keep the EARLIEST copy of a duplicate (ties: first listed) so the 24h clock starts at the original request
    w_["K"] = (f'=IF(B{r}="","",IF(SUMPRODUCT(({D_("J")}=J{r})*({D_("H")}<H{r}))'
               f'+SUMPRODUCT(($J$2:J{r}=J{r})*($H$2:H{r}=H{r}))-1>0,"DUPLICATE",'
               f'IF(AND(ISNUMBER(SEARCH("@",E{r})),ISNUMBER(SEARCH(".",E{r})),LEN(F{r})=10,ISNUMBER(VALUE(F{r}))),"OK","BAD CONTACT")))')
    w_["L"] = f'=IF(B{r}="",0,IF(K{r}="DUPLICATE",0,1))'
    hit = lambda code: f'COUNTIFS({LP},F{r},{LC},B{r},{LO},"{code}")+COUNTIFS({LP},F{r},{LC},"",{LO},"{code}")'
    w_["M"] = f'=IF(L{r}=1,IF({hit("*BOOKED")}>0,1,0),0)'
    w_["N"] = f'=IF(L{r}=1,IF({hit("NOT_INTERESTED")}>0,1,0),0)'
    w_["O"] = f'=IF(L{r}=1,IF(COUNTIFS({LP},F{r},{LO},"STOP")+COUNTIFS({LP},F{r},{LO},"WRONG_NUMBER")>0,1,0),0)'
    pst = f'IFERROR(INDEX({PCOL("D")},MATCH(B{r},{PA},0)),"")'
    w_["P"] = (f'=IF(L{r}=1,IF({pst}="NO SLOTS",1,IF(AND({pst}<>"REOPENED",'
               f'COUNTIFS({LC},B{r},{LO},"NO_SLOTS")>0),1,0)),0)')
    w_["Q"] = f'=IF(L{r}=1,COUNTIFS({LP},F{r},{LTRY},1),0)'
    w_["R"] = f'=IF(Q{r}>0,_xlfn.MAXIFS({LT},{LP},F{r},{LTRY},1),0)'
    w_["S"] = f'=IF(Q{r}>0,IFERROR(INDEX({LO},MATCH(F{r}&"|"&Q{r},{LKEY},0)),""),"")'
    w_["T"] = f'=IF(S{r}="",0,IFERROR(VLOOKUP(S{r},{OUTR},4,FALSE),0))'
    w_["U"] = f'=IF(Q{r}=0,0,R{r}+T{r}/1440)'
    w_["V"] = f'=IF(L{r}=1,H{r}+{SET["sla"]}/24,"")'
    w_["W"] = f'=IF(L{r}=1,(V{r}-{CK})*24,"")'
    w_["X"] = (f'=IF(AND(L{r}=1,K{r}="OK",M{r}=0,N{r}=0,O{r}=0,P{r}=0,{CK}<V{r},Q{r}<{SET["maxtry"]}),1,0)')
    w_["Y"] = (f'=IF(AND(X{r}=1,{CK}>=H{r},{CK}>=U{r},{TOD}>={mins("win0")},{TOD}<{mins("win1")}),1,0)')
    w_["Z"] = (f'=IF(B{r}="","",IF({TOD}<{mins("shy0")},"Ram",IF({TOD}>={mins("ram1")},"Shyam",'
               f'IF(MOD(IFERROR(VALUE(RIGHT(F{r},2)),0),2)=0,"Ram","Shyam"))))')
    w_["AA"] = f'=IF(X{r}=1,COUNTIFS({D_("F")},F{r},{D_("X")},1),0)'
    w_["AB"] = f'=IF(X{r}=1,COUNTIFS($F$2:F{r},F{r},$X$2:X{r},1),0)'
    w_["AC"] = f'=IF(L{r}=1,IF(IFERROR(INDEX({PCOL("C")},MATCH(B{r},{PA},0)),"")="YES","YES","NO"),"")'
    w_["AD"] = f'=IF(L{r}=1,IFERROR(INDEX({PCOL("B")},MATCH(B{r},{PA},0)),""),"")'
    w_["AE"] = f'=IF(X{r}=1,IFERROR(INDEX(SETTINGS!$B$30:$D$32,MATCH(I{r},SETTINGS!$A$30:$A$32,0),MIN(Q{r},2)+1),SETTINGS!$B$32),0)'
    w_["AF"] = f'=IF(L{r}=1,IFERROR(VALUE(RIGHT(G{r},1)),1),0)'
    w_["AG"] = (f'=IF(AND(Y{r}=1,AB{r}=1),ROUND(AA{r}*AE{r}*IF(AF{r}>1,{SET["rbst"]},1)'
                f'*IF(W{r}<={SET["urgent"]},{SET["ubst"]},1),6),0)')
    w_["AH"] = (f'=IF(AG{r}>0,COUNTIFS({D_("Z")},Z{r},{D_("AG")},">"&AG{r})'
                f'+COUNTIFS($Z$2:Z{r},Z{r},$AG$2:AG{r},AG{r}),"")')
    w_["AI"] = f'=IF(AH{r}<>"",Z{r}&"|"&AH{r},"")'
    w_["AJ"] = f'=IF(X{r}=1,F{r}&"|"&AB{r},"")'
    w_["AK"] = f'=IF(L{r}=1,IF(COUNTIFS($F$2:F{r},F{r},$L$2:L{r},1)=1,1,0),0)'
    first = lambda c: f'_xlfn.MINIFS({LT},{LP},F{r},{LC},{c},{LO},"*BOOKED")'
    w_["AM"] = (f'=IF(M{r}=1,IF({first(f"B{r}")}=0,{first(chr(34)*2)},IF({first(chr(34)*2)}=0,{first(f"B{r}")},'
                f'MIN({first(f"B{r}")},{first(chr(34)*2)}))),"")')
    w_["AN"] = f'=IF(M{r}=1,(AM{r}-H{r})*24,"")'
    w_["AO"] = f'=IF(AND(M{r}=1,AN{r}<>""),IF(AN{r}<={SET["sla"]},1,0),0)'
    w_["AP"] = f'=IF(M{r}=1,IF({hit("BOOKED")}>0,1,0),0)'
    w_["AQ"] = f'=IF(B{r}="","",LEFT(D{r},FIND(" ",D{r}&" ")-1))'
    w_["AL"] = (f'=IF(B{r}="","",IF(K{r}="DUPLICATE","Duplicate (removed)",IF(AO{r}=1,"Scheduled within 24h",'
                f'IF(M{r}=1,"Scheduled late",IF(K{r}="BAD CONTACT","Bad contact details",'
                f'IF(O{r}=1,"Stopped / wrong number",IF(N{r}=1,"Not interested",IF(P{r}=1,"Partner has no slots",'
                f'IF({CK}>=V{r},IF(Q{r}=0,"Missed deadline: never tried","Missed deadline: tried, not booked"),'
                f'IF(Q{r}>={SET["maxtry"]},"No reply after max tries","Open"))))))))))')
    for col, fml in w_.items():
        c = ws[f"{col}{r}"]
        c.value = fml
        c.font = f_norm
    for col in ("H", "R", "U", "V", "AM"):
        ws[f"{col}{r}"].number_format = "yyyy-mm-dd hh:mm"
    ws[f"F{r}"].number_format = "@"
ws.freeze_panes = "C2"


# ---------------------------------------------------------------------------------- RAM / SHYAM lists (interns)
def msg_formula(code, name, intern, company, link):
    part = lambda k: f'INDEX(LISTS!$B$14:$B$19,MATCH({k},LISTS!$A$14:$A$19,0))'
    txt = f'{part(code)}&IF({link}="",{part(chr(34) + "YES" + chr(34))},{part(chr(34) + "NO" + chr(34))})'
    return (f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({txt},"{{NAME}}",{name}),"{{INTERN}}",{intern}),'
            f'"{{COMPANY}}",{company}),"{{LINK}}",{link})')


def build_list(sheet, who):
    ws = S[sheet]
    shift = ("ram0", "ram1") if who == "Ram" else ("shy0", "shy1")
    ws["A1"], ws["A1"].font = f"{who.upper()}: your list. Start at the top. Do what column G says.", f_title
    ws["B2"] = who                          # formulas read the intern's name from B2
    ws["C2"], ws["D2"] = "Time now", f"={CK}"
    ws["D2"].number_format = "ddd dd mmm  hh:mm"
    ws["D2"].alignment = Alignment(horizontal="left")
    ws["G2"] = f'=C3&" people waiting"'
    ws["C3"] = f'=SUMPRODUCT((DESK!$Z$2:$Z${R}=$B$2)*(DESK!$AI$2:$AI${R}<>""))'
    ws["N2"] = f"={SET['urgent']}"          # local copy: Sheets conditional formats cannot point at other tabs
    for c_ in ("B2", "C2", "G2"):
        ws[c_].font = f_bold
    ws["D2"].font = f_norm
    ws["C3"].font = Font(name=F, size=10, color="FFFFFF")
    ws["N2"].font = Font(name=F, size=8, color="FFFFFF")
    ws["A3"] = (f'=IF({TOD}<{mins(shift[0])},"Your shift has not started. Nothing to do yet.",'
                f'IF({TOD}>={mins(shift[1])},"Your shift is over. Thank you!",'
                f'IF($C$3=0,"Nothing waiting right now. Check LOG column J for lines marked in orange.","")))')
    ws["A3"].font = Font(name=F, size=11, bold=True, color="C00000")
    # look-up box: someone replied on WhatsApp or called back
    ws["A4"] = "Reply or call-back? Phone:"
    ws["A4"].font = f_bold
    ws.merge_cells("A4:C4")
    ws["D4"].font, ws["D4"].fill, ws["D4"].border = f_input, fill_in, box
    ws["D4"].number_format = "@"
    lk = lambda k: f'IFERROR(MATCH($D$4&"|{k}",DESK!$AJ$2:$AJ${R},0),"")'
    ws["E4"] = (f'=IF($D$4="","",IF({lk(1)}="","Not on the open list (already done, or wrong number).",'
                f'INDEX(DESK!$B$2:$B${R},{lk(1)})&" "&INDEX(DESK!$G$2:$G${R},{lk(1)})'
                f'&IF({lk(2)}="","","; "&INDEX(DESK!$B$2:$B${R},{lk(2)})&" "&INDEX(DESK!$G$2:$G${R},{lk(2)}))'
                f'&IF({lk(3)}="","","; "&INDEX(DESK!$B$2:$B${R},{lk(3)})&" "&INDEX(DESK!$G$2:$G${R},{lk(3)}))'
                f'&IF(INDEX(DESK!$AC$2:$AC${R},{lk(1)})="YES"," : they said yes? Book it, log BOOKED."," : send the link, log SENT_LINK.")))')
    ws["E4"].font = f_bold
    ws.merge_cells("E4:I4")
    for j, col in enumerate(("J", "K", "L"), start=1):
        ws[f"{col}4"] = (f'=IF({lk(j)}="","",HYPERLINK(INDEX(DESK!$AD$2:$AD${R},{lk(j)}),'
                         f'"Calendar: "&INDEX(DESK!$B$2:$B${R},{lk(j)})))')
        ws[f"{col}4"].font = f_link
    header(ws, 6, ["#", "Candidate", "Phone", "Interviews to book", "Hours left", "Try", "WHAT TO DO",
                   "Call", "WhatsApp", "Calendar 1", "Calendar 2", "Calendar 3", "Message (copy if WhatsApp opens empty)",
                   "key", "row", "msg"],
           widths=[4, 16, 13, 30, 7, 5, 58, 15, 16, 14, 14, 14, 70, 4, 4, 4])
    for k in range(1, NQ + 1):
        r = 6 + k
        ws[f"A{r}"] = k
        ws[f"N{r}"] = f'=$B$2&"|"&A{r}'
        ws[f"O{r}"] = f'=IFERROR(MATCH(N{r},DESK!$AI$2:$AI${R},0),"")'
        ix = lambda col: f'INDEX(DESK!${col}$2:${col}${R},O{r})'
        m2 = f'MATCH($C{r}&"|2",DESK!$AJ$2:$AJ${R},0)'
        m3 = f'MATCH($C{r}&"|3",DESK!$AJ$2:$AJ${R},0)'
        ws[f"B{r}"] = f'=IF(O{r}="","",{ix("D")})'
        ws[f"C{r}"] = f'=IF(O{r}="","",{ix("F")})'
        ws[f"D{r}"] = (f'=IF(O{r}="","",{ix("B")}&" "&{ix("G")}'
                       f'&IFERROR("; "&INDEX(DESK!$B$2:$B${R},{m2})&" "&INDEX(DESK!$G$2:$G${R},{m2}),"")'
                       f'&IFERROR("; "&INDEX(DESK!$B$2:$B${R},{m3})&" "&INDEX(DESK!$G$2:$G${R},{m3}),""))')
        ws[f"E{r}"] = f'=IF(O{r}="","",ROUND({ix("W")},1))'
        ws[f"F{r}"] = f'=IF(O{r}="","",{ix("Q")}+1)'
        last = f'({ix("S")}="SENT_LINK")'
        may = f'({ix("AC")}="YES")'
        ws[f"P{r}"] = f'=IF(O{r}="","",IF({last},"CHK",IF(F{r}=1,"T1",IF(F{r}>={SET["maxtry"]},"T3","T2"))))'
        ws[f"G{r}"] = (f'=IF(O{r}="","",IF(F{r}>={SET["maxtry"]},"LAST TRY. ","")&IF({last},'
                       f'"Open Calendar 1. Already booked? Log SELF_BOOKED. Not booked? Call, send the WhatsApp, log what happened.",'
                       f'IF({may},"Call. They say yes? Open Calendar 1, book the first time that suits them, log BOOKED. '
                       f'No answer? Send the WhatsApp, log NO_ANSWER.",'
                       f'"Call. They say yes? Send the WhatsApp (it has the booking link), log SENT_LINK. '
                       f'No answer? Send the WhatsApp, log NO_ANSWER.")))')
        cnt = f'{ix("AA")}'
        company = f'{ix("B")}&IF({cnt}>1," and "&({cnt}-1)&" more","")'
        link = f'IF({may},"",{ix("AD")})'
        ws[f"M{r}"] = f'=IF(O{r}="","",{msg_formula(f"P{r}", ix("AQ"), "$B$2", company, link)})'
        url = f'"https://wa.me/91"&C{r}&"?text="&_xlfn.ENCODEURL(M{r})'
        ws[f"I{r}"] = (f'=IF(O{r}="","",IF(LEN({url})<=255,HYPERLINK({url},"Open WhatsApp"),'
                       f'HYPERLINK("https://wa.me/91"&C{r},"Open WhatsApp, paste column M")))')
        ws[f"H{r}"] = f'=IF(O{r}="","",HYPERLINK("tel:+91"&C{r},"Call "&C{r}))'
        ws[f"J{r}"] = f'=IF(O{r}="","",HYPERLINK({ix("AD")},{ix("B")}))'
        for col, m in (("K", m2), ("L", m3)):
            ws[f"{col}{r}"] = (f'=IF(O{r}="","",IFERROR(HYPERLINK(INDEX(DESK!$AD$2:$AD${R},{m}),'
                               f'INDEX(DESK!$B$2:$B${R},{m})),""))')
        for col in "ABCDEFGHIJKLM":
            ws[f"{col}{r}"].font = f_norm
            ws[f"{col}{r}"].alignment = wrap
        for col in "HIJKL":
            ws[f"{col}{r}"].font = f_link
        ws[f"G{r}"].font = f_bold
        ws[f"C{r}"].number_format = "@"
        ws.row_dimensions[r].height = 44
    for col in "NOP":
        ws.column_dimensions[col].hidden = True
    ws.conditional_formatting.add(f"A7:M{6 + NQ}", FormulaRule(formula=['AND(ISNUMBER($E7),$E7<=$N$2)'], fill=fill_urgent))
    ws.freeze_panes = "B7"


build_list("RAM", "Ram")
build_list("SHYAM", "Shyam")

# ---------------------------------------------------------------------------------- DASHBOARD (lead)
ws = S["DASHBOARD"]
ws["A1"], ws["A1"].font = "DASHBOARD (lead). Everything updates live.", f_title
ws.column_dimensions["A"].width = 62
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 70
DK = lambda c: f"DESK!${c}$2:${c}${R}"


def kv(row, label, formula, fmt=None, note=None, bold=False):
    ws.cell(row, 1, label).font = f_bold if bold else f_norm
    c = ws.cell(row, 2, formula)
    c.font = f_bold
    if fmt:
        c.number_format = fmt
    if note:
        ws.cell(row, 3, note).font = f_note
        ws.cell(row, 3).alignment = wrap


ws["A3"] = "Scheduled within 24h (share of all unique requests)"
ws["A3"].font = Font(name=F, size=12, bold=True)
ws["B3"] = f'=IF(B6=0,0,SUM({DK("AO")})/B6)'
ws["B3"].font, ws["B3"].number_format = f_big, "0%"
ws["C3"] = f'="Target "&TEXT({SET["target"]},"0%")&IF(B3>={SET["target"]},":  ON TARGET",":  below target")'
ws["C3"].font = Font(name=F, size=12, bold=True)
ws.row_dimensions[3].height = 30
section(ws, 5, "Today", 3)
kv(6, "Unique requests (the denominator for every share)", f"=SUM({DK('L')})")
kv(7, "Duplicate rows removed", f'=COUNTIF({DK("K")},"DUPLICATE")')
kv(8, "People", f"=SUM({DK('AK')})")
kv(9, "Scheduled within 24h", f"=SUM({DK('AO')})")
kv(10, "   booked by an intern", f"=SUMPRODUCT({DK('AO')},{DK('AP')})")
kv(11, "   booked by the candidate", "=B9-B10")
kv(12, "Open now", f'=COUNTIF({DK("AL")},"Open")')
kv(13, "At risk (open, deadline within the urgent window)", f'=COUNTIFS({DK("X")},1,{DK("W")},"<="&{SET["urgent"]})')
kv(14, "Tries made", f"=SUM({LTRY})")
kv(15, "Requests booked per try (measured)", '=IF(B14=0,"-",B10/B14)', "0.00")
kv(16, "Intern minutes per interview booked", f'=IF(B10=0,"-",B14*{SET["mintry"]}/B10)', "0.0")
section(ws, 18, "Why requests are not scheduled", 3)
states = ["Open", "Missed deadline: never tried", "Missed deadline: tried, not booked", "No reply after max tries",
          "Bad contact details", "Stopped / wrong number", "Not interested", "Partner has no slots", "Scheduled late"]
for k, s_ in enumerate(states):
    kv(19 + k, s_, f'=COUNTIF({DK("AL")},"{s_}")')
    ws.cell(19 + k, 3, f'=IF($B$6=0,"",B{19 + k}/$B$6)').number_format = "0%"
    ws.cell(19 + k, 3).font = f_norm
    ws.cell(19 + k, 3).alignment = Alignment(horizontal="left")

section(ws, 29, "Capacity: can today's staff reach the target?", 3)
dl = f"MOD(MAX({DK('V')}),1)"
kv(30, "Intern minutes before the deadline (Ram + Shyam)",
   f'=IF(B6=0,0,(MAX(0,MIN({SET["ram1"]},{dl})-{SET["ram0"]})+MAX(0,MIN({SET["shy1"]},{dl})-{SET["shy0"]}))*1440)', "0",
   "Assumes one evening batch, worked the next day.")
kv(31, "Tries the interns can make", f'=ROUNDDOWN(B30/{SET["mintry"]},0)', "0")
kv(32, "Requests booked per try (measured after 30 tries, else planning value)",
   f'=IF(B14>=30,B15,{SET["pbook"]})', "0.00")
kv(33, "Share candidates book themselves (planning value)", f'={SET["pself"]}', "0%",
    "From SETTINGS. Measured the next morning as row 11 / row 6.")
kv(34, "Requests needed for the target", f'=ROUNDUP({SET["target"]}*B6,0)', "0")
kv(35, "   of which interns must book", "=MAX(0,B34-ROUND(B33*B6,0))", "0")
kv(36, "Tries needed", '=IF(B32=0,"-",ROUNDUP(B35/B32,0))', "0")
kv(37, "Most tries that can be useful (people x max tries)", f'=B8*{SET["maxtry"]}', "0",
   "Optimistic: assumes later tries book as often as early ones.")
kv(38, "Interns needed (same shifts as today)", '=IF(OR(B30=0,NOT(ISNUMBER(B36))),"-",ROUNDUP(B36/(B31/2),0))', "0")
ws["A39"], ws["A39"].font = "Verdict", f_bold
ws["B39"] = ('=IF(NOT(ISNUMBER(B36)),"-",IF(B36>B37,"Adding interns cannot reach the target: raise reply rates '
             '(automated message, upstream re-confirmation) or agree a different target.",'
             'IF(B36<=B31,"Today\'s staff can reach the target.","Needs about "&B38&" interns on today\'s shifts.")))')
ws["B39"].font, ws["B39"].alignment = f_bold, wrap
ws.merge_cells("B39:C39")
ws.row_dimensions[39].height = 30

section(ws, 41, "Cost: is extra capacity worth paying for?", 3)
kv(42, "Cost of one intern-day (Rs)", f'={SET["stip"]}/{SET["days"]}', "#,##0")
kv(43, "Cost per try (Rs)", f'=B42/(B31/2)', "#,##0.0")
kv(44, "Cost per interview booked by an intern (Rs)", '=IF(ISNUMBER(B32),B43/B32,"-")', "#,##0")
kv(45, "Expected value of one scheduled interview to Weekday (Rs)",
   f'={SET["fee"]}*{SET["ctc"]}*{SET["hire"]}', "#,##0", "Fee x CTC x hires per interview. All three from SETTINGS.")
kv(46, "Value / cost", '=IF(ISNUMBER(B44),B45/B44,"-")', "0", "How many times over a booked interview repays its cost.")
kv(47, "Gap to target today (requests)", f'=MAX(0,B34-B9)', "0")
kv(48, "Revenue at risk from the gap, per day (Rs)", f'=B47*{SET["lost"]}*B45', "#,##0",
   "Only the share of late requests that are never interviewed counts as lost.")
kv(49, "Extra interns to close the gap, cost per day (Rs)",
   '=IF(ISNUMBER(B38),MAX(0,B38-2)*B42,"not reachable by staffing")', "#,##0")
kv(50, "Automated evening WhatsApp to everyone, 2 messages, cost per day (Rs)",
   f'=B8*2*{SET["wamsg"]}', "#,##0", "Message fees only. A WhatsApp provider adds a monthly platform fee.")
ws["A52"] = ("Read this as a rule, not a forecast: while one booked interview is worth many times what it costs, adding capacity "
             "pays. Capacity stops helping once everyone has had their tries (row 36 above row 37); after that only reply "
             "rates move the number.")
ws["A52"].font, ws["A52"].alignment = f_note, wrap
ws.merge_cells("A52:C52")
ws.row_dimensions[52].height = 40
section(ws, 54, "Do later tries still work? (the number that decides whether more interns can reach the target)", 3)
header(ws, 55, ["Try number", "Tries made", "Ended in BOOKED or SENT_LINK"])
for j in (1, 2, 3):
    r = 55 + j
    ws.cell(r, 1, f"Try {j}").font = f_norm
    ws.cell(r, 2, f'=COUNTIF({LKEY},"*|{j}")').font = f_norm
    ws.cell(r, 3, f'=IF(B{r}=0,"-",(COUNTIFS({LKEY},"*|{j}",{LO},"BOOKED")+COUNTIFS({LKEY},"*|{j}",{LO},"SENT_LINK"))/B{r})')
    ws.cell(r, 3).number_format = "0%"
    ws.cell(r, 3).font = f_bold
    ws.cell(r, 3).alignment = Alignment(horizontal="left")
ws["A59"] = ("Row 38 assumes every try books as often as the average. If try 3 books far less than try 1, extra interns mostly "
             "buy useless third tries, and only reply rates (the evening message, upstream re-confirmation) move the number.")
ws["A59"].font, ws["A59"].alignment = f_note, wrap
ws.merge_cells("A59:C59")
ws.row_dimensions[59].height = 30

# ---------------------------------------------------------------------------------- START HERE (interns)
ws = S["START HERE"]
ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 92
ws["B1"], ws["B1"].font = "How to work the list", f_title
steps = [
    ("1", "Open your tab", "RAM or SHYAM. Start at row 1."),
    ("2", "Do what column G says", "Click Call. Then book, or send the WhatsApp, as column G says."),
    ("3", "Log it", "Go to LOG. Add ONE line: time (Ctrl+Shift+;), your name, the phone, the outcome."),
    ("4", "Repeat", "Go back to your tab. The row you finished goes away and the next one moves up."),
]
for i, (n_, a, b) in enumerate(steps, start=3):
    ws.cell(i, 1, n_).font = Font(name=F, size=14, bold=True, color="1F3864")
    ws.cell(i, 2, a).font = Font(name=F, size=11, bold=True)
    ws.cell(i, 3, b).font = Font(name=F, size=11)
    ws.row_dimensions[i].height = 22
ws["B8"], ws["B8"].font = "Which outcome?", f_title
header(ws, 9, ["", "Outcome", "When"])
for i, (code, when, *_r) in enumerate(outcomes, start=10):
    ws.cell(i, 2, code).font = f_bold
    ws.cell(i, 3, when).font = f_norm
    for c_ in (2, 3):
        ws.cell(i, c_).border = box
ws["B20"], ws["B20"].font = "Good to know", f_title
notes = [
    ("Two companies, two answers", "Yes to one, no to the other? Write two lines: BOOKED and NOT_INTERESTED, each with the company in column E."),
    ("Someone replies later", "Type their phone in the yellow box at the top of your tab (\"Reply or call-back? Phone\"). It shows what to book."),
    ("Orange cell in LOG", "Something is wrong with that line. Column J says what. Fix it."),
    ("Yellow row", "Less than 4 hours left. These rows are already at the top."),
    ("Never", "Contact anyone before 10:00 or after 21:00. Promise anything about the job or the salary. Change other tabs."),
    ("Not sure?", "Write it in Notes and ask your lead."),
]
for i, (a, b) in enumerate(notes, start=21):
    ws.cell(i, 2, a).font = f_bold
    ws.cell(i, 2).alignment = wrap
    c = ws.cell(i, 3, b)
    c.font, c.alignment = f_norm, wrap
    ws.row_dimensions[i].height = 28

# ---------------------------------------------------------------------------------- finish
for s_ in wb.worksheets:
    s_.sheet_view.showGridLines = s_.title not in ("START HERE", "DASHBOARD")
for _n in ("DESK", "LISTS"):
    S[_n].sheet_state = "hidden"
for _n in ("START HERE", "RAM", "SHYAM", "LOG"):
    S[_n].sheet_properties.tabColor = "70AD47"
for _n in ("DASHBOARD", "PARTNERS", "INPUT", "SETTINGS"):
    S[_n].sheet_properties.tabColor = "4472C4"
wb.active = 0
wb.calculation.fullCalcOnLoad = True      # the file carries no cached values; Excel computes everything on open
wb.save(OUT)
print("saved", OUT, "| raw rows", len(rows), "| partners", NPART, "| YES-consent", sum(v == "YES" for v in consent.values()),
      "| ptab", ptab)
